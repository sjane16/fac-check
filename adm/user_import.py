"""
Import users from Excel (User Management > Import users).

  1. import_template  - download an .xlsx with the right columns (no password column)
  2. import_preview   - upload the filled file; every row is checked with the same rules
                        as the signup form and shown back before anything is created
  3. import_confirm   - create the valid rows. Each person gets a random temporary
                        password by email and must choose their own at first login.
"""
import io
import json
import secrets
import string
import time

from django.conf import settings
from django.contrib.auth.models import User
from django.core.exceptions import ValidationError
from django.core.mail import EmailMessage, get_connection
from django.core.validators import validate_email
from django.db import transaction
from django.http import HttpResponse, JsonResponse
from django.views.decorators.http import require_POST
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from security.audit import log_event
from user.models import UserProfile
from user.validators import validate_name

# Each account means a slow password hash plus one Gmail email (Gmail allows ~500/day),
# so keep one import to a size that finishes in a single request.
MAX_IMPORT_ROWS = 100
MAX_FILE_BYTES = 2 * 1024 * 1024
SESSION_KEY = 'um_import'
SESSION_MAX_AGE = 30 * 60

COLUMNS = ['First Name', 'Last Name', 'Email', 'Role', 'Department', 'Phone']
ROLES = ['student', 'faculty', 'staff', 'med_personnel']
# Ways people write the MED role in a spreadsheet
ROLE_ALIASES = {'med personnel': 'med_personnel', 'med': 'med_personnel', 'med-personnel': 'med_personnel'}
DEPARTMENTS = dict(UserProfile.DEPARTMENT_CHOICES + UserProfile.OFFICE_CHOICES)  # code -> full name (colleges and offices)


def _error(message, status=400):
    return JsonResponse({'ok': False, 'error': message}, status=status)


# ---------- 1. template ----------

def import_template(request):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Users'
    ws.append(COLUMNS)
    header_fill = PatternFill('solid', fgColor='2563EB')
    for cell in ws[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = header_fill
    for col, width in zip('ABCDEF', (18, 18, 32, 12, 16, 16)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = 'A2'
    # Phone as text so Excel keeps the leading 0
    for row in range(2, MAX_IMPORT_ROWS + 2):
        ws[f'F{row}'].number_format = '@'

    role_dv = DataValidation(type='list', formula1='"student,faculty,staff,med_personnel"', allow_blank=False)
    role_dv.error, role_dv.errorTitle = 'Choose student, faculty, staff or med_personnel.', 'Invalid role'
    dept_dv = DataValidation(type='list', formula1='"' + ','.join(DEPARTMENTS) + '"', allow_blank=False)
    dept_dv.error, dept_dv.errorTitle = 'Choose a department code from the list.', 'Invalid department'
    ws.add_data_validation(role_dv)
    ws.add_data_validation(dept_dv)
    role_dv.add(f'D2:D{MAX_IMPORT_ROWS + 1}')
    dept_dv.add(f'E2:E{MAX_IMPORT_ROWS + 1}')

    help_ws = wb.create_sheet('Instructions')
    lines = [
        ('How to fill in this file', True),
        ('Fill in one person per row on the "Users" sheet. Do not rename the columns.', False),
        (f'Up to {MAX_IMPORT_ROWS} people per file.', False),
        ('', False),
        ('First Name / Last Name - letters, spaces, hyphens, apostrophes and periods only.', False),
        ('Email - students and faculty must use their @udm.edu.ph email; staff and MED personnel may use any email.', False),
        ('Role - student, faculty, staff or med_personnel (MED Personnel; "MED" also works). Admins are added from the User Management page.', False),
        ('Department - one of the codes below: a college for students and faculty, an office or area for staff.', False),
        ('Phone - 11 digits, e.g. 09171234567.', False),
        ('', False),
        ('There is no password column: Fac-Check emails each person a temporary password,', False),
        ('and they choose their own password the first time they log in.', False),
        ('', False),
        ('Department codes', True),
        ('Colleges (students and faculty)', True),
    ] + [(f'{code} - {name}', False) for code, name in UserProfile.DEPARTMENT_CHOICES] + [
        ('', False),
        ('Offices and areas (staff)', True),
    ] + [(f'{code} - {name}', False) for code, name in UserProfile.OFFICE_CHOICES]
    for text, bold in lines:
        help_ws.append([text])
        if bold:
            help_ws.cell(row=help_ws.max_row, column=1).font = Font(bold=True, size=12)
    help_ws.column_dimensions['A'].width = 95

    buffer = io.BytesIO()
    wb.save(buffer)
    response = HttpResponse(buffer.getvalue(),
                            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="fac-check-user-import-template.xlsx"'
    return response


# ---------- 2. preview ----------

def _text(value):
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()

def _clean_phone(value):
    phone = ''.join(ch for ch in _text(value) if ch.isdigit())
    # Excel often turns 09171234567 into the number 9171234567
    if len(phone) == 10 and phone.startswith('9'):
        phone = '0' + phone
    return phone

def _clean_department(value):
    raw = _text(value)
    for code, name in DEPARTMENTS.items():
        if code and raw.lower() in (code.lower(), name.lower()):
            return code
    return raw

def _check_row(data, seen_emails, existing_emails):
    errors = []
    for key, label in (('first_name', 'First name'), ('last_name', 'Last name')):
        try:
            data[key] = validate_name(data[key], field_label=label)
        except ValidationError as e:
            errors.extend(e.messages)

    email = data['email'].lower()
    data['email'] = email
    try:
        validate_email(email)
    except ValidationError:
        errors.append("Email is not a valid email address.")
    else:
        if data['role'] in ('student', 'faculty') and not email.endswith('@udm.edu.ph'):
            errors.append("Students and faculty must use their @udm.edu.ph email.")
        if email in existing_emails:
            errors.append("An account with this email already exists.")
        elif email in seen_emails:
            errors.append("This email appears more than once in the file.")
    seen_emails.add(email)

    if data['role'] not in ROLES:
        errors.append("Role must be student, faculty, staff or med_personnel.")
    if data['department'] not in DEPARTMENTS or not data['department']:
        errors.append("Department must be one of the codes on the Instructions sheet.")
    elif data['role'] in ROLES and data['department'] not in dict(UserProfile.department_choices_for(data['role'])):
        errors.append("Staff need an office or area code; students and faculty need a college code (see the Instructions sheet).")
    if not (len(data['phone']) == 11 and data['phone'].isdigit()):
        errors.append("Phone number must be exactly 11 digits, numbers only.")
    return errors


@require_POST
def import_preview(request):
    upload = request.FILES.get('file')
    if not upload:
        return _error("Choose an Excel file to upload.")
    if not upload.name.lower().endswith('.xlsx'):
        return _error("Upload an Excel file (.xlsx). Download the template if you're not sure.")
    if upload.size > MAX_FILE_BYTES:
        return _error("This file is too large (max 2 MB).")

    try:
        wb = load_workbook(upload, read_only=True, data_only=True)
    except Exception:
        return _error("This file couldn't be read. Make sure it's an Excel (.xlsx) file based on the template.")
    ws = wb['Users'] if 'Users' in wb.sheetnames else wb.worksheets[0]
    rows = ws.iter_rows(values_only=True)

    header = [_text(h).lower() for h in (next(rows, None) or [])]
    wanted = {c.lower(): c for c in COLUMNS}
    positions = {}
    for idx, name in enumerate(header):
        if name in wanted:
            positions[name] = idx
    missing = [wanted[c] for c in wanted if c not in positions]
    if missing:
        wb.close()
        return _error("Missing column(s): " + ", ".join(missing) + ". Use the template so the columns match.")

    existing = set(User.objects.values_list('email', flat=True)) | set(User.objects.values_list('username', flat=True))
    existing = {e.lower() for e in existing if e}
    seen, results = set(), []
    for line_no, row in enumerate(rows, start=2):
        values = {name: (row[idx] if idx < len(row) else None) for name, idx in positions.items()}
        if not any(_text(v) for v in values.values()):
            continue  # blank line
        if len(results) >= MAX_IMPORT_ROWS:
            wb.close()
            return _error(f"This file has more than {MAX_IMPORT_ROWS} people. Split it into smaller files.")
        data = {
            'first_name': _text(values['first name']),
            'last_name': _text(values['last name']),
            'email': _text(values['email']),
            'role': ROLE_ALIASES.get(_text(values['role']).lower(), _text(values['role']).lower()),
            'department': _clean_department(values['department']),
            'phone': _clean_phone(values['phone']),
        }
        errors = _check_row(data, seen, existing)
        results.append({'row': line_no, **data, 'department_name': DEPARTMENTS.get(data['department'], ''), 'errors': errors})
    wb.close()

    if not results:
        return _error("The file has no people in it. Fill in the Users sheet and upload it again.")

    valid = [r for r in results if not r['errors']]
    token = secrets.token_urlsafe(16)
    request.session[SESSION_KEY] = {
        'token': token, 'at': time.time(),
        'rows': [{k: r[k] for k in ('row', 'first_name', 'last_name', 'email', 'role', 'department', 'phone')} for r in valid],
    }
    return JsonResponse({'ok': True, 'token': token, 'rows': results,
                         'valid_count': len(valid), 'error_count': len(results) - len(valid)})


# ---------- 3. create ----------

def temporary_password(length=12):
    """Random password that meets the Fac-Check rules (upper, lower, number, symbol)."""
    symbols = '!@#$%^&*?'
    pools = [string.ascii_uppercase, string.ascii_lowercase, string.digits, symbols]
    chars = [secrets.choice(pool) for pool in pools]
    everything = string.ascii_letters + string.digits + symbols
    chars += [secrets.choice(everything) for _ in range(length - len(chars))]
    secrets.SystemRandom().shuffle(chars)
    return ''.join(chars)


def _welcome_email(row, password, connection, login_url):
    body = (
        f"Hi {row['first_name']},\n\n"
        "An administrator created a Fac-Check account for you.\n\n"
        f"Email: {row['email']}\n"
        f"Temporary password: {password}\n\n"
        f"Log in at {login_url}. You'll receive a verification code by email, then you'll be asked "
        "to choose your own password.\n\n"
        "Please don't share this email. If you weren't expecting it, you can ignore it."
    )
    EmailMessage('Your Fac-Check account', body, settings.DEFAULT_FROM_EMAIL, [row['email']],
                 connection=connection).send(fail_silently=False)


@require_POST
def import_confirm(request):
    try:
        body = json.loads(request.body or b'{}')
    except ValueError:
        return _error("Invalid request.")
    data = request.session.get(SESSION_KEY)
    if not data or body.get('token') != data.get('token') or time.time() - data.get('at', 0) > SESSION_MAX_AGE:
        return _error("This preview expired. Upload the file again.", status=409)
    request.session.pop(SESSION_KEY, None)

    login_url = request.build_absolute_uri('/auth/login/')
    created, failed = [], []
    connection = get_connection(fail_silently=False)
    try:
        connection.open()
    except Exception:
        return _error("Couldn't connect to the email server, so no accounts were created. Try again later.", status=503)

    try:
        for row in data['rows']:
            email = row['email']
            if User.objects.filter(username__iexact=email).exists() or User.objects.filter(email__iexact=email).exists():
                failed.append({'email': email, 'reason': 'An account with this email was created in the meantime.'})
                continue
            password = temporary_password()
            try:
                # The account only exists if its email was sent, so nobody ends up
                # with an account they can't get into
                with transaction.atomic():
                    user = User.objects.create_user(
                        username=email, email=email, password=password,
                        first_name=row['first_name'], last_name=row['last_name'], is_active=True,
                    )
                    UserProfile.objects.update_or_create(user=user, defaults={
                        'role': row['role'], 'department': row['department'], 'phone': row['phone'],
                        'approval_status': UserProfile.APPROVAL_APPROVED,
                        'must_change_password': True,
                    })
                    _welcome_email(row, password, connection, login_url)
                created.append(email)
            except Exception:
                failed.append({'email': email, 'reason': "Couldn't send the welcome email, so the account wasn't created."})
    finally:
        connection.close()

    log_event(request, 'users_imported', user=request.user, created=len(created), failed=len(failed))
    return JsonResponse({'ok': True, 'created': created, 'failed': failed})
